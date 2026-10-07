"""Build the bundled paper UPD status 2 template for the client package.

The layout follows the recommended FNS UPD form and filling guidance in
letter 21.10.2013 No. MMV-20-3/96@ (nalog.gov.ru/rn77/about_fts/about_nalog/4319219/).
Status 2 is a primary transfer document, not a VAT invoice. The DOCX is for
printing and signatures; it is not the regulated electronic XML format.
"""

import base64
import hashlib
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "resources" / "document-templates"
NAME = "client_upd_status2_v1.docx"

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(29.7), Cm(21)
section.left_margin = section.right_margin = Cm(1.2)
section.top_margin = section.bottom_margin = Cm(1.1)
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(8)
normal.paragraph_format.space_after = Pt(2)


def para(text="", bold=False, size=8, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    return p


def cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(7.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


para("УНИВЕРСАЛЬНЫЙ ПЕРЕДАТОЧНЫЙ ДОКУМЕНТ", True, 12, WD_ALIGN_PARAGRAPH.CENTER)
para("Статус: 2 — передаточный документ (акт); не является счётом-фактурой", True, 9)
para("Документ № {{upd.number}} от {{upd.date}}    Исправление № — от —")
para("Продавец (исполнитель): {{company.name}}")
para("Адрес: {{company.legal_address}}    ИНН: {{company.inn}}    {{company.kpp_line}}")
para("Покупатель (заказчик): {{client.name}}")
para("Адрес: {{client.legal_address}}    ИНН: {{client.inn}}    {{client.kpp_line}}")
para("Грузоотправитель и грузополучатель: — (документ об оказании услуг)")
para("Валюта: российский рубль, код 643")

headers = [
    "№ п/п", "Наименование товара, описание работ или услуг", "Ед. изм.",
    "Кол-во", "Цена, ₽", "Стоимость без налога, ₽", "Ставка НДС / сумма налога", "Стоимость всего, ₽",
]
table = doc.add_table(rows=3, cols=len(headers))
table.style = "Table Grid"
widths = [Cm(1), Cm(10), Cm(1.6), Cm(1.6), Cm(2.7), Cm(3.4), Cm(3.4), Cm(3.4)]
for row in table.rows:
    for i, cell in enumerate(row.cells):
        cell.width = widths[i]
for i, value in enumerate(headers):
    cell_text(table.cell(0, i), value, True)
for i, value in enumerate([
    "1", "{{service.description}}", "усл.", "1", "{{service.amount}}",
    "{{service.amount}}", "Без НДС / —", "{{service.amount}}",
]):
    cell_text(table.cell(1, i), value)
for i, value in enumerate(["", "Всего к передаче", "", "", "", "{{service.amount}}", "—", "{{service.amount}}"]):
    cell_text(table.cell(2, i), value, True)

para("Всего к оплате: {{service.amount_words}}. НДС не предъявляется.", True)
para("Основание передачи (сдачи): договор субаренды транспортного средства с экипажем № {{contract.number}} от {{contract.date_plain}}")
para("Услуги оказаны по рейсу № {{trip.business_id}}. Иные сведения: —")

sign = doc.add_table(rows=2, cols=2)
sign.style = "Table Grid"
cell_text(sign.cell(0, 0), "Услуги оказал / передал")
cell_text(sign.cell(0, 1), "Услуги принял / получил")
cell_text(sign.cell(1, 0), "{{company.director_position}}  __________________  {{company.director_full_name}}\nДата передачи: {{upd.date}}\n{{company.name}}")
cell_text(sign.cell(1, 1), "__________________  ______________________________\nДолжность / подпись / расшифровка\nДата получения: «____» ______________ 20____ г.\n{{client.name}}")
para("Экземпляр для подписания сторонами. Дата приёмки и подпись получателя заполняются при фактическом принятии услуг.", size=7)

tmp = OUT / NAME
doc.save(tmp)
source = tmp.read_bytes()
buffer = io.BytesIO()
with zipfile.ZipFile(io.BytesIO(source)) as original, zipfile.ZipFile(buffer, "w") as stable:
    for member in original.infolist():
        fixed = zipfile.ZipInfo(member.filename, (1980, 1, 1, 0, 0, 0))
        fixed.compress_type = member.compress_type
        fixed.external_attr = member.external_attr
        stable.writestr(fixed, original.read(member.filename))
data = buffer.getvalue()
tmp.write_bytes(data)
digest = hashlib.sha256(data).hexdigest()
encoded = base64.b64encode(data).decode("ascii")
for index, offset in enumerate(range(0, len(encoded), 60000), 1):
    (OUT / f"{NAME}.b64.part{index:03d}").write_text(encoded[offset:offset + 60000] + "\n", encoding="ascii")
print(f"{NAME} SHA256={digest} parts={(len(encoded) + 59999) // 60000}")
